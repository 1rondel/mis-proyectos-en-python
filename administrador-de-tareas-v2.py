from openpyxl import load_workbook, Workbook
import os
def cargar_tarea():
    #funcion que permita cargar tarea(tareas, estados) de una hoja de excel (en este caso) y que tambien sea capaz de crear una hoja de excel si no existe
    if os.path.exists('tareas.xlsx'):
        wb = load_workbook('tareas.xlsx')
        hoja = wb.active
        tareas = []
        estados = []
        for fila in hoja.iter_rows(min_row=2, values_only=True):
            descripcion = fila[0]
            estado_texto = fila[1]
            if descripcion is not None: #hacer que se salte lasfilas vacias
             tareas.append(descripcion)
             estados.append(estado_texto == 'completado')
            
    else:
             respuesta = input("Archivo inexistente, desea crear uno nuevo, s/n").lower()
             if respuesta == 's':
                print('Creando archivo de guardado...')
                wb = Workbook()
                hoja = wb.active
                hoja.append(['Descripcion', 'Estados'])
                wb.save('tareas.xlsx')
                print('Archivo creado con exito')
                return [], []
             else:
                print('Continuando sin guardado')
                return [], []
    return tareas, estados

def guardar_tarea(tareas, estados):
    #guardar tarea en su hoja correspondiente
    wb = Workbook()
    hoja = wb.active
    hoja.append(['Descripcion', 'Estados'])
    for tarea, estado in zip(tareas, estados):
        estado_texto = "completado" if estado else "pendiente"
        hoja.append([tarea, estado_texto])
    wb.save('tareas.xlsx')



def mostrar_menu():
     print('Menu de opciones')
     print('1. Agregar tareas')
     print('2. Ver tareas con estados')
     print('3. Marcar tarea como completada')
     print('4. Salir del programa')

def agregar_tareas(tareas, estados):
    #debes de crear una funcion para agregar tareas
    descripcion = input('Describa su tarea')
    tareas.append(descripcion)
    estados.append(False)
    guardar_tarea(tareas, estados)
    print("Tarea agregada exitosamente")

def ver_tareas_estados(tareas, estados):
    if not tareas:
        print('No hay tareas')
        return
    else:
     for i in range(len(tareas)):
      estado = '[X]' if estados[i] else '[ ]'
      print(f'{i+1}. {estado} {tareas[i]}')
      
    # para ver tareas junto con su estado
    
def marcar_completado(tareas, estados):
    #debes de hacer una funcion que marque como completado la tarea seleccionada
    if len(tareas) == 0:
            print('No hay Tareas')
            return
    else:
        for i in range(len(tareas)):
         estado = '[X]' if estados[i] else '[ ]'
         print(f'{i+1}. {estado} {tareas[i]}')
    try:
     numero = int(input('ingresa el numero de la tarea que haz completado'))
    
     indice = numero - 1
     if 0 <= indice < len(tareas):
               estados[indice] = True
               guardar_tarea(tareas, estados)
               print('Felicidades la tarea fue completada')
     else:
              print('numero invalido')
    except ValueError:
       print('Debes de ingresar un numero valido')
       

#------administrador-------
tareas, estados = cargar_tarea()
while True:
    try:
     mostrar_menu()
     funcion = int(input('Marque el numero de opcion'))
     if funcion == 1:
        agregar_tareas(tareas, estados)
     elif funcion == 2:
        ver_tareas_estados(tareas, estados)
     elif funcion == 3:
        marcar_completado(tareas, estados)
     elif funcion == 4:
        guardar_tarea(tareas, estados)
        print('Finalizando programa')
        break
     else:
         print('ingrese numero valido entre el 1 y 4') 
    except ValueError:
        print('Ingrese un numero valido')
        continue